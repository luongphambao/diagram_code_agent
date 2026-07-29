# -*- coding: utf-8 -*-
"""Cập nhật lại phần mục lục (TOC) đã cache trong file .docx.

Trường TOC của Word chỉ được tính lại khi mở bằng Word. Để file bàn giao đọc
đúng ngay cả trong trình xem PDF/preview, ta ghi lại phần kết quả cache của
trường bằng số trang thật, lấy từ bookmark của bản PDF do LibreOffice kết xuất.

Chạy 2 vòng vì việc đổi số dòng mục lục làm thay đổi phân trang.
"""
import os
import subprocess
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import docx_edit as de

SOFFICE = "/usr/bin/soffice"


def to_pdf(docx_path, outdir):
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                    docx_path, "--outdir", outdir],
                   check=True, capture_output=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def pdf_outline(pdf_path):
    import warnings
    warnings.filterwarnings("ignore")
    import pypdf
    r = pypdf.PdfReader(pdf_path)
    out = []

    def walk(items, depth=0):
        for it in items:
            if isinstance(it, list):
                walk(it, depth + 1)
            else:
                try:
                    pg = r.get_destination_page_number(it) + 1
                except Exception:
                    pg = 1
                out.append((depth, it.title.strip(), pg))

    walk(r.outline)
    return out


def _run(p, text=None, tab=False, style=None):
    r = p.add_run(text or "")
    if style:
        try:
            r.style = style
        except KeyError:
            pass
    if tab:
        r._r.append(OxmlElement("w:tab"))
    return r


def rebuild_toc(docx_path, entries, save_as=None):
    """entries: list[(level0, title, page)] — level0 = 0/1/2 tương ứng toc 1/2/3."""
    doc = Document(docx_path)
    bs = de.blocks(doc)
    toc_idx = [i for i, b in enumerate(bs)
               if hasattr(b, "style") and (b.style.name or "").startswith("toc ")]
    if not toc_idx:
        raise RuntimeError("Không tìm thấy đoạn mục lục nào trong tài liệu.")

    first, last = bs[toc_idx[0]], bs[toc_idx[-1]]

    # giữ lại 3 run mở trường (begin / instrText / separate) và run đóng trường
    open_runs, close_run = [], None
    for r in first._p.findall(qn("w:r")):
        fc = r.find(qn("w:fldChar"))
        if (fc is not None and fc.get(qn("w:fldCharType")) in ("begin", "separate")) \
                or r.find(qn("w:instrText")) is not None:
            open_runs.append(r)
    for r in last._p.iter(qn("w:r")):
        fc = r.find(qn("w:fldChar"))
        if fc is not None and fc.get(qn("w:fldCharType")) == "end":
            close_run = r
    for r in open_runs:
        r.getparent().remove(r)
    if close_run is not None:
        close_run.getparent().remove(close_run)

    # mượn pPr (tab stop, dot leader, khoảng cách) của mục lục gốc theo từng cấp
    import copy as _copy
    exemplar = {}
    for i in toc_idx:
        b = bs[i]
        name = b.style.name
        if name not in exemplar:
            pPr = b._p.find(qn("w:pPr"))
            if pPr is not None:
                exemplar[name] = _copy.deepcopy(pPr)

    anchor = first._p.getprevious()
    parent = first._p.getparent()
    for i in toc_idx:
        el = de._element(bs[i])
        el.getparent().remove(el)

    new_ps = []
    for lvl, title, page in entries:
        style = f"toc {min(lvl, 8) + 1}"
        p = doc.add_paragraph()
        try:
            p.style = doc.styles[style]
        except KeyError:
            style = "toc 1"
            p.style = doc.styles[style]
        if style in exemplar:
            old = p._p.find(qn("w:pPr"))
            if old is not None:
                p._p.remove(old)
            p._p.insert(0, _copy.deepcopy(exemplar[style]))
        # tách "3.2.1 Description" -> số hiệu | tiêu đề để khớp 2 tab stop của style
        import re as _re
        m = _re.match(r"^((?:\d+(?:\.\d+)*)|(?:Appendix\s+[A-Z]\.))\s+(.+)$", title)
        if m:
            _run(p, m.group(1))
            _run(p, tab=True)
            _run(p, m.group(2))
        else:
            _run(p, title)
        _run(p, tab=True)
        _run(p, str(page))
        new_ps.append(p)

    if new_ps:
        for r in reversed(open_runs):
            new_ps[0]._p.insert(1, r)
        if close_run is not None:
            new_ps[-1]._p.append(close_run)

    cursor = anchor
    for p in new_ps:
        p._p.getparent().remove(p._p)
        if cursor is None:
            parent.insert(0, p._p)
        else:
            cursor.addnext(p._p)
        cursor = p._p

    doc.save(save_as or docx_path)
    return len(new_ps)


def refresh(docx_path, workdir="out", passes=2):
    for k in range(passes):
        pdf = to_pdf(docx_path, workdir)
        entries = pdf_outline(pdf)
        n = rebuild_toc(docx_path, entries)
        print(f"  vòng {k + 1}: {n} dòng mục lục, {len(entries)} heading")
    pdf = to_pdf(docx_path, workdir)
    return pdf


if __name__ == "__main__":
    print(refresh(sys.argv[1] if len(sys.argv) > 1
                  else "out/BRD_Diagram_Code_Agent_v1.0.docx"))
