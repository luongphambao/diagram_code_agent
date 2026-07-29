"""BRD eval runner — deterministic, no agent/LLM, CI-safe.

cd backend
uv run python -m evals.brd.run_eval
uv run python -m evals.brd.run_eval --gate
uv run python -m evals.brd.run_eval --update-baseline
"""

from __future__ import annotations

import argparse
import io
import sys
import tempfile
from pathlib import Path

_BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_BACKEND / "src"))
sys.path.insert(0, str(_BACKEND / "src" / "domain" / "brd"))
sys.path.insert(0, str(_BACKEND))

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

import brd_docx as bd  # noqa: E402
from brd_validate import validate_brd  # noqa: E402
from evals._core import (  # noqa: E402
    compare_to_baseline,
    load_cases,
    print_table,
    run_all_sync,
    write_baseline,
    write_results,
)
from evals.brd.judge import METRIC_KEYS, score_brd  # noqa: E402

_DIR = Path(__file__).resolve().parent
_DATASET = _DIR / "dataset"
_RESULTS = _DIR / "results.json"
_BASELINE = _DIR / "baseline.json"

_COLUMNS = [
    ("case_id", "case"),
    ("scores.ok", "ok"),
    ("scores.recall", "recall"),
]

# Standalone multilevel-heading numbering for the eval fixtures — an unused id
# (90), matching the same scheme tests/test_brd_docx.py uses, so a bare
# Document() has real numPr-backed headings and neither
# heading_missing_numbering (validate case) nor the "heading mất đánh số tự
# động" semantic-preservation check fire spuriously.
_HEADING_NUM_ID = 90


def _register_numbering(doc) -> None:
    numbering_el = doc.part.numbering_part.element
    abstract = f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="{_HEADING_NUM_ID}">
      <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
      <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
      <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3."/></w:lvl>
    </w:abstractNum>"""
    num = f'<w:num {nsdecls("w")} w:numId="{_HEADING_NUM_ID}"><w:abstractNumId w:val="{_HEADING_NUM_ID}"/></w:num>'
    numbering_el.append(parse_xml(abstract))
    numbering_el.append(parse_xml(num))


def _heading(doc, text: str, level: int):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = level - 1
    numPr.get_or_add_numId().val = _HEADING_NUM_ID
    return p


def _build_doc(headings: list[dict]) -> "Document":
    doc = Document()
    _register_numbering(doc)
    for h in headings:
        _heading(doc, h["title"], h["level"])
        for p in h.get("paragraphs", []):
            doc.add_paragraph(p)
    return doc


def _to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _insert_heading_after(doc, after_title: str, title: str, level: int, paragraphs: list[str]) -> None:
    """Append a heading + its paragraphs, then splice them right after the
    `after_title` paragraph — mirrors op_insert_section's element-move trick,
    without going through the tool-facing op API (this is testing INDEXING
    stability, not the op itself)."""
    target = next(p for p in doc.paragraphs if p.text.strip() == after_title)
    new_els = [_heading(doc, title, level)._p] + [doc.add_paragraph(t)._p for t in paragraphs]
    for el in new_els:
        el.getparent().remove(el)
    cursor = target._p
    for el in new_els:
        cursor.addnext(el)
        cursor = el


def _run_id_stability(case: dict) -> dict:
    doc = _build_doc(case["headings"])
    _insert_heading_after(
        doc,
        case["insert_after_title"],
        case["insert_title"],
        case["insert_level"],
        case.get("insert_paragraphs", []),
    )
    _, sections = bd.index_document(doc)
    return {"after_ids": {s.section_id for s in sections}}


def _run_checksum_guard(case: dict) -> dict:
    doc = _build_doc(case["headings"])
    _, sections = bd.index_document(doc)
    sec = bd.resolve_section_ref(sections, case["section"])
    before_bytes = _to_bytes(doc)

    succeeded, _applied, failed = bd.apply_ops(
        doc,
        [{"op": case["op"], "section": case["section"], "expect": "0" * 16, "content": case["content"]}],
    )
    unchanged_on_reject = before_bytes == _to_bytes(doc)
    rejected_stale = len(succeeded) == 0 and len(failed) == 1

    succeeded2, _applied2, _failed2 = bd.apply_ops(
        doc,
        [
            {
                "op": case["op"],
                "section": case["section"],
                "expect": sec.own_checksum,
                "content": case["content"],
            }
        ],
    )
    accepted_when_correct = len(succeeded2) == 1
    return {
        "rejected_stale": rejected_stale,
        "unchanged_on_reject": unchanged_on_reject,
        "accepted_when_correct": accepted_when_correct,
    }


def _run_semantic_preservation(case: dict) -> dict:
    doc = _build_doc(case["headings"])
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / f"{case['id']}.docx"
        doc.save(path)
        _, sections = bd.index_document(path)
        sec = bd.resolve_section_ref(sections, case["section"])

        result = bd.apply_brd_ops(
            path,
            [
                {
                    "op": "rename_heading",
                    "section": case["section"],
                    "expect": sec.own_checksum,
                    "title": case["new_title"],
                }
            ],
            revisions_dir=Path(td) / "revs",
        )
        _, after = bd.index_document(path)
        return {"ok": result.ok, "reverted": result.reverted, "after_ids": {s.section_id for s in after}}


def _run_validate(case: dict) -> dict:
    doc = _build_doc(case["headings"])
    report = validate_brd(doc)
    return {"errors": report["errors"]}


_RUNNERS = {
    "id_stability": _run_id_stability,
    "checksum_guard": _run_checksum_guard,
    "semantic_preservation": _run_semantic_preservation,
    "validate": _run_validate,
}


def _run_one(case: dict) -> dict:
    runner = _RUNNERS.get(case["kind"])
    if runner is None:
        raise ValueError(f"unknown BRD eval case kind: {case['kind']!r}")
    result = runner(case)
    return {"case_id": case.get("id", "unknown"), "scores": score_brd(result, case)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", default=None)
    ap.add_argument("--gate", action="store_true")
    ap.add_argument("--update-baseline", action="store_true")
    args = ap.parse_args()

    cases = load_cases(_DATASET, args.case)
    if not cases:
        print(f"No brd cases in {_DATASET}")
        sys.exit(1)

    results = run_all_sync(cases, _run_one)
    print_table(results, _COLUMNS)
    write_results(_RESULTS, results)

    if args.update_baseline:
        write_baseline(_BASELINE, results, METRIC_KEYS)
        print(f"Baseline updated: {_BASELINE}")
        return

    passed, regressions = compare_to_baseline(results, _BASELINE, METRIC_KEYS)
    if regressions:
        print("REGRESSIONS:")
        for r in regressions:
            print(f"  {r['metric']}: {r['baseline']} -> {r['current']} (drop {r['drop']})")
    else:
        print("No regression vs baseline." if _BASELINE.exists() else "No baseline yet.")
    if args.gate and not passed:
        sys.exit(1)


if __name__ == "__main__":
    main()
