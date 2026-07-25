"""Render the WBS Excel deliverable's key sheets to PNG, so the proposal deck can embed
the REAL ``wbs_filled.xlsx`` look (brand styling, conditional formatting, the client's
actual live-formula numbers) instead of a plainer re-derived python-pptx table for the
same data — the client-visible artifact and the deck illustration of it never diverge.

Pipeline (one workbook copy, one soffice call, one pdftoppm call):
  1. Load ``wbs_filled.xlsx``, drop every sheet except the 3 worth showing a client
     (:data:`SHEET_KINDS` — "1. Effort" / "2. WBS" / "3. Delivery Plan"; the "0. How to
     use" cover and "4. Master Data" ratios sheet are internal, not proposal content).
  2. Set each kept sheet's print area to its used range + landscape fit-to-one-page, so
     the PDF export doesn't split a sheet across multiple pages or crop content.
  3. ``soffice --headless --convert-to pdf`` the trimmed workbook -> one PDF, one page
     per remaining sheet, in sheet order.
  4. ``pdftoppm -png`` the PDF -> one PNG per page.
  5. Write ``wbs_sheet_images.json`` — ``{"effort": "<png path>", "wbs": "...",
     "delivery": "..."}`` — the manifest ``ppt_reporting`` checks (mirroring the existing
     ``tech_icons.json`` / ``icon_plan.json`` "prefer the resolved asset over the native
     render" pattern already used for the tech-stack slide).

Best-effort, NEVER blocking: if LibreOffice/poppler aren't installed (e.g. running
outside the Docker image that bundles them) or anything else goes wrong, every entry
point here returns ``{}`` / does nothing instead of raising — callers keep the native
python-pptx table/gantt renderers as the fallback, so a workspace outside the container
is never blocked on this feature.
"""

from __future__ import annotations

import json
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image, ImageChops

logger = logging.getLogger(__name__)

SHEET_IMAGES_NAME = "wbs_sheet_images.json"

# Sheet name (in wbs_filled.xlsx, see wbs_excel._KEEP_SHEETS) -> manifest key. Kept in
# workbook order so the PDF's page order matches this dict's iteration order.
SHEET_KINDS: dict[str, str] = {
    "1. Effort": "effort",
    "2. WBS": "wbs",
    "3. Delivery Plan": "delivery",
}

_CONVERT_TIMEOUT_S = 60


def soffice_available() -> bool:
    return shutil.which("soffice") is not None and shutil.which("pdftoppm") is not None


def _fit_to_one_landscape_page(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = ws.dimensions


def _prepare_render_copy(xlsx_path: Path, dest_path: Path) -> list[str]:
    """Clone ``xlsx_path`` keeping only :data:`SHEET_KINDS`' sheets, each set to print as
    one landscape page. Returns the kept sheet names in their final (PDF page) order."""
    wb = load_workbook(xlsx_path)
    for name in list(wb.sheetnames):
        if name not in SHEET_KINDS:
            del wb[name]
    kept = [name for name in SHEET_KINDS if name in wb.sheetnames]
    for name in kept:
        _fit_to_one_landscape_page(wb[name])
    wb.active = 0
    wb.save(dest_path)
    return kept


def render_wbs_sheets(
    xlsx_path: Path,
    workspace: Path,
    *,
    dpi: int = 150,
    timeout: int = _CONVERT_TIMEOUT_S,
) -> dict[str, str]:
    """Render :data:`SHEET_KINDS`' sheets of ``xlsx_path`` to PNG under ``workspace``,
    write ``wbs_sheet_images.json``, and return the same ``{kind: filename}`` mapping.

    Returns ``{}`` (and writes nothing) if LibreOffice/poppler are unavailable, the
    workbook can't be read, or the conversion fails for any reason — see the module
    docstring on why this never raises.
    """
    if not Path(xlsx_path).exists():
        return {}
    if not soffice_available():
        logger.info("soffice/pdftoppm not found on PATH — skipping WBS sheet-image render.")
        return {}

    try:
        with tempfile.TemporaryDirectory(prefix="wbs_render_") as tmp:
            tmp_dir = Path(tmp)
            tmp_xlsx = tmp_dir / "wbs_render.xlsx"
            kept_sheets = _prepare_render_copy(Path(xlsx_path), tmp_xlsx)
            if not kept_sheets:
                return {}

            # A per-call UserInstallation profile avoids soffice's shared-profile lock,
            # which would otherwise serialize (or fail) concurrent renders from different
            # per-thread workspaces (this is a multi-tenant agent backend — see backends.py
            # §4.10 per-thread isolation).
            subprocess.run(
                [
                    "soffice",
                    "--headless",
                    "--norestore",
                    f"-env:UserInstallation=file://{tmp_dir / 'lo_profile'}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_dir),
                    str(tmp_xlsx),
                ],
                check=True,
                capture_output=True,
                timeout=timeout,
            )
            tmp_pdf = tmp_xlsx.with_suffix(".pdf")
            if not tmp_pdf.exists():
                logger.warning("soffice did not produce a PDF for %s", xlsx_path)
                return {}

            page_prefix = tmp_dir / "page"
            subprocess.run(
                ["pdftoppm", "-png", "-r", str(dpi), str(tmp_pdf), str(page_prefix)],
                check=True,
                capture_output=True,
                timeout=timeout,
            )

            manifest: dict[str, str] = {}
            for i, sheet_name in enumerate(kept_sheets, start=1):
                kind = SHEET_KINDS[sheet_name]
                # pdftoppm zero-pads the page number only when there are >=10 pages; with
                # <=3 sheets it's always a bare "-1"/"-2"/"-3" suffix.
                src = tmp_dir / f"page-{i}.png"
                if not src.exists():
                    continue
                dest_name = f"wbs_sheet_{kind}.png"
                dest = Path(workspace) / dest_name
                shutil.copyfile(src, dest)
                manifest[kind] = dest_name

            if manifest:
                (Path(workspace) / SHEET_IMAGES_NAME).write_text(
                    json.dumps(manifest, indent=2), encoding="utf-8"
                )
            return manifest
    except Exception as exc:  # noqa: BLE001 — rendering is best-effort, never blocking
        logger.warning("WBS sheet-image render failed (%s) — falling back to native tables.", exc)
        return {}


def load_wbs_sheet_images(workspace: Path) -> dict[str, str]:
    """Read the manifest written by :func:`render_wbs_sheets` ``{}`` if absent/unreadable."""
    path = Path(workspace) / SHEET_IMAGES_NAME
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def resolve_wbs_sheet_image(workspace: Path, kind: str) -> Optional[Path]:
    """Absolute path to the rendered PNG for ``kind`` ("effort"/"wbs"/"delivery"), or
    ``None`` if not rendered / the file has since been removed."""
    manifest = load_wbs_sheet_images(workspace)
    ref = manifest.get(kind)
    if not ref:
        return None
    p = Path(ref)
    p = p if p.is_absolute() else Path(workspace) / p
    return p if p.exists() else None
