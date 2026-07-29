"""Tests for tools/analysis/brd_tools.py (docs/plans/2026-07-29-brd-agent.md §E2)."""

from __future__ import annotations

import contextvars
import json

import backends
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from tools import (
    draft_brd_outline,
    draft_section_content,
    import_brd_docx,
    inspect_brd_template,
    load_brd_context,
    read_brd_outline,
    validate_brd,
)
from tools.schemas.brd import Block, OutlineItem


def _use_workspace(monkeypatch, tmp_path) -> None:
    tmp_path.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(
        backends,
        "_current_workspace",
        contextvars.ContextVar("current_workspace", default=tmp_path),
    )


_NUM_ID = 93


def _numbered_doc():
    doc = Document()
    numbering_el = doc.part.numbering_part.element
    abstract = f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="{_NUM_ID}">
      <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
      <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
    </w:abstractNum>"""
    num = f'<w:num {nsdecls("w")} w:numId="{_NUM_ID}"><w:abstractNumId w:val="{_NUM_ID}"/></w:num>'
    numbering_el.append(parse_xml(abstract))
    numbering_el.append(parse_xml(num))
    p = doc.add_paragraph("INTRODUCTION", style="Heading 1")
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = _NUM_ID
    p2 = doc.add_paragraph("Purpose", style="Heading 2")
    numPr2 = p2._p.get_or_add_pPr().get_or_add_numPr()
    numPr2.get_or_add_ilvl().val = 1
    numPr2.get_or_add_numId().val = _NUM_ID
    doc.add_paragraph("Muc tieu tai lieu.")
    return doc


def test_load_brd_context_warns_when_workspace_is_empty(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    out = load_brd_context.func()
    assert "CẢNH BÁO" in out


def test_load_brd_context_reads_real_artifacts(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    (tmp_path / "out.png").write_bytes(b"x")
    (tmp_path / "diagram_brief.json").write_text(
        json.dumps({"objective": "Ban ve", "functional_requirements": ["A", "B"]}), encoding="utf-8"
    )
    out = load_brd_context.func()
    assert "Ban ve" in out
    assert "Functional requirements: 2" in out


def test_inspect_brd_template_indexes_the_real_template(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    out = inspect_brd_template.func()
    assert "introduction" in out
    assert (tmp_path / "template_map.json").exists()
    cached = json.loads((tmp_path / "template_map.json").read_text(encoding="utf-8"))
    assert any(s["section_id"] == "introduction/purpose" for s in cached)


def test_inspect_brd_template_missing_file_is_a_plain_message(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    out = inspect_brd_template.func(template="does-not-exist.docx")
    assert "Không tìm thấy template" in out


def test_draft_brd_outline_and_section_content_round_trip(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    items = [
        OutlineItem(
            section_id="introduction/purpose", title="Purpose", level=2, source="manual", status="fill"
        ),
        OutlineItem(
            section_id="introduction/scope",
            title="Scope",
            level=2,
            source="manual",
            status="skip",
            notes="chua co blueprint",
        ),
    ]
    out = draft_brd_outline.func(items=items)
    assert "2 mục" in out
    assert (tmp_path / "brd_outline_draft.json").exists()

    content = [Block(type="p", text="Tai lieu nay dung de..."), Block(type="bullet", text="Muc tieu 1")]
    out2 = draft_section_content.func(section_id="introduction/purpose", content=content)
    assert "introduction/purpose" in out2
    saved = json.loads((tmp_path / "brd_sections" / "introduction__purpose.json").read_text(encoding="utf-8"))
    assert saved["section_id"] == "introduction/purpose"
    assert len(saved["content"]) == 2


def test_read_brd_outline_before_any_doc_exists_says_so(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    out = read_brd_outline.func(section="")
    assert "import_brd_docx" in out


def test_read_brd_outline_shows_the_draft_when_only_a_draft_exists(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    draft_brd_outline.func(
        items=[OutlineItem(section_id="introduction/purpose", title="Purpose", level=2, source="manual")]
    )
    out = read_brd_outline.func(section="")
    assert "Nháp outline" in out
    assert "introduction/purpose" in out


def test_read_brd_outline_whole_doc_and_drilldown(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    _numbered_doc().save(tmp_path / "out.brd.docx")

    whole = read_brd_outline.func(section="")
    assert "introduction/purpose" in whole
    assert "chk=" in whole

    drilled = read_brd_outline.func(section="purpose")
    assert "Mục 'introduction/purpose'" in drilled
    assert "own_checksum=" in drilled


def test_read_brd_outline_ambiguous_or_unknown_section_is_a_plain_error(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    _numbered_doc().save(tmp_path / "out.brd.docx")
    out = read_brd_outline.func(section="does-not-exist")
    assert out.startswith("✗")


def test_validate_brd_before_doc_exists(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    assert "Chưa có" in validate_brd.func()


def test_validate_brd_reports_findings_and_writes_manifest(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    doc = _numbered_doc()
    doc.add_paragraph("TODO: dien noi dung")
    doc.save(tmp_path / "out.brd.docx")
    out = validate_brd.func()
    assert "lỗi" in out
    assert (tmp_path / "brd_validation.json").exists()


def test_import_brd_docx_from_a_workspace_relative_filename(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    _numbered_doc().save(tmp_path / "client_brd.docx")
    out = import_brd_docx.func(source="client_brd.docx")
    assert "introduction/purpose" in out
    assert (tmp_path / "out.brd.docx").exists()


def test_import_brd_docx_rejects_a_path_escaping_the_workspace(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    out = import_brd_docx.func(source="../../etc/passwd")
    assert "Không tìm thấy" in out
    assert not (tmp_path / "out.brd.docx").exists()


def test_import_brd_docx_rejects_a_corrupt_file(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    (tmp_path / "broken.docx").write_bytes(b"not a real docx")
    out = import_brd_docx.func(source="broken.docx")
    assert "không phải .docx hợp lệ" in out
    assert not (tmp_path / "out.brd.docx").exists()


def test_import_brd_docx_from_an_uploaded_file_id(monkeypatch, tmp_path):
    _use_workspace(monkeypatch, tmp_path)
    from backends import AGENT_SPACE

    uploads_dir = AGENT_SPACE / "uploads"
    uploads_dir.mkdir(parents=True, exist_ok=True)
    file_id = "abc123def456"
    dest = uploads_dir / f"{file_id}_client_brd.docx"
    _numbered_doc().save(dest)
    try:
        out = import_brd_docx.func(source=file_id)
        assert "introduction/purpose" in out
        assert (tmp_path / "out.brd.docx").exists()
    finally:
        dest.unlink(missing_ok=True)
