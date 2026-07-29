"""Tests for domain/brd/brd_validate.py (docs/plans/2026-07-29-brd-agent.md §D)."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

import brd_validate as bv

_NUM_ID = 92


def _register_numbering(doc) -> None:
    numbering_el = doc.part.numbering_part.element
    abstract = f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="{_NUM_ID}">
      <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
      <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
      <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3."/></w:lvl>
    </w:abstractNum>"""
    num = f'<w:num {nsdecls("w")} w:numId="{_NUM_ID}"><w:abstractNumId w:val="{_NUM_ID}"/></w:num>'
    for xml in (abstract, num):
        numbering_el.append(parse_xml(xml))


def _heading(doc, text: str, level: int, *, numbered: bool = True):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if numbered:
        numPr = p._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = level - 1
        numPr.get_or_add_numId().val = _NUM_ID
    return p


def _clean_doc():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "INTRODUCTION", 1)
    _heading(doc, "Purpose", 2)
    doc.add_paragraph("Muc tieu that su cua tai lieu nay, du dai de qua nguong 20 ky tu.")
    _heading(doc, "FUNCTIONAL REQUIREMENTS", 1)
    _heading(doc, "FR01 – Login", 2)
    _heading(doc, "Description", 3)
    doc.add_paragraph("Nguoi dung dang nhap.")
    _heading(doc, "Interface requirements", 3)
    doc.add_paragraph("Form dang nhap.")
    _heading(doc, "Data requirements", 3)
    doc.add_paragraph("Bang users.")
    return doc


def _bytes_of(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_clean_document_has_no_errors():
    result = bv.validate_brd(io.BytesIO(_bytes_of(_clean_doc())))
    assert result["ok"] is True
    assert result["error_count"] == 0
    assert result["section_count"] == 7


def test_placeholder_residue_is_an_error():
    doc = _clean_doc()
    doc.add_paragraph("TODO: dien vao sau")
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    codes = [e["code"] for e in result["errors"]]
    assert "placeholder_residue" in codes


def test_heading_without_numbering_is_an_error():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "INTRODUCTION", 1, numbered=False)
    doc.add_paragraph("noi dung")
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    assert any(e["code"] == "heading_missing_numbering" for e in result["errors"])


def test_heading_level_jump_is_an_error():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "INTRODUCTION", 1)
    doc.add_paragraph("noi dung")
    _heading(doc, "Deep detail", 3)  # jumps from level 1 straight to level 3
    doc.add_paragraph("chi tiet")
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    assert any(e["code"] == "heading_level_jump" for e in result["errors"])


def test_table_column_mismatch_is_an_error():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "BUDGET", 1)
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[1]._tr.append(
        tbl.rows[1].cells[0]._tc.makeelement(qn("w:tc"), {})
    )  # ragged row: 3 cells vs 2 header cols
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    assert any(e["code"] == "table_column_mismatch" for e in result["errors"])


def test_fr_section_missing_required_children_is_an_error():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "FUNCTIONAL REQUIREMENTS", 1)
    _heading(doc, "FR01 – Login", 2)
    _heading(doc, "Description", 3)
    doc.add_paragraph("chi co description, thieu 2 muc con con lai")
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    findings = [e for e in result["errors"] if e["code"] == "fr_missing_slot"]
    assert findings and "interface-requirements" in findings[0]["message"]


def test_ambiguous_sibling_title_is_a_warning():
    doc = Document()
    _register_numbering(doc)
    _heading(doc, "A", 1)
    _heading(doc, "Description", 2)
    doc.add_paragraph("x")
    _heading(doc, "Description", 2)
    doc.add_paragraph("y")
    result = bv.validate_brd(io.BytesIO(_bytes_of(doc)))
    assert any(w["code"] == "ambiguous_anchor" for w in result["warnings"])


def test_thin_fill_section_warning_uses_outline_param():
    doc = _clean_doc()
    result = bv.validate_brd(
        io.BytesIO(_bytes_of(doc)),
        outline={"functional-requirements/fr01/interface-requirements": {"status": "fill"}},
    )
    assert any(w["code"] == "thin_fill_section" for w in result["warnings"])
