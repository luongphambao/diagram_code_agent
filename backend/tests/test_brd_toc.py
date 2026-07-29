"""Tests for domain/brd/brd_toc.py (docs/plans/2026-07-29-brd-agent.md §E1)."""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import brd_toc


def _add_toc_field_paragraph(doc, *, title: str = "Cu 1", page: str = "1"):
    try:
        doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["Normal"]
    except ValueError:
        pass  # already exists
    p = doc.add_paragraph(style="toc 1")

    def _fld(kind):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)

    _fld("begin")
    r = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    r._r.append(instr)
    _fld("separate")
    p.add_run(f"{title}\t{page}")
    _fld("end")
    return p


def test_soffice_available_returns_a_bool():
    assert isinstance(brd_toc.soffice_available(), bool)


def test_mark_toc_dirty_counts_and_sets_the_flag():
    doc = Document()
    _add_toc_field_paragraph(doc)
    n = brd_toc._mark_toc_dirty(doc)
    assert n == 1
    fld = next(doc.element.body.iter(qn("w:fldChar")))
    assert fld.get(qn("w:dirty")) == "true"


def test_mark_toc_dirty_on_a_doc_without_toc_is_a_noop():
    doc = Document()
    doc.add_paragraph("no toc here")
    assert brd_toc._mark_toc_dirty(doc) == 0


def test_refresh_toc_falls_back_to_dirty_only_when_soffice_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(brd_toc.shutil, "which", lambda name: None)
    doc = Document()
    _add_toc_field_paragraph(doc)
    path = tmp_path / "doc.docx"
    doc.save(path)

    ok = brd_toc.refresh_toc(path)
    assert ok is False

    reopened = Document(path)
    fld = next(reopened.element.body.iter(qn("w:fldChar")))
    assert fld.get(qn("w:dirty")) == "true"


def test_rebuild_toc_rewrites_entries_and_preserves_the_field(tmp_path):
    doc = Document()
    _add_toc_field_paragraph(doc, title="Cu cu", page="99")
    path = tmp_path / "doc.docx"
    doc.save(path)

    n = brd_toc.rebuild_toc(path, [(0, "1 Gioi thieu", 5), (0, "2 Pham vi", 8)])
    assert n == 2

    reopened = Document(path)
    toc_paragraphs = [p for p in reopened.paragraphs if (p.style.name or "") == "toc 1"]
    assert len(toc_paragraphs) == 2
    assert "Gioi thieu" in toc_paragraphs[0].text and "5" in toc_paragraphs[0].text
    assert "Pham vi" in toc_paragraphs[1].text and "8" in toc_paragraphs[1].text

    field_kinds = [fc.get(qn("w:fldCharType")) for fc in reopened.element.body.iter(qn("w:fldChar"))]
    assert field_kinds == ["begin", "separate", "end"], "the field must survive the rewrite intact"


@pytest.mark.skipif(not brd_toc.soffice_available(), reason="soffice not installed in this environment")
def test_refresh_toc_end_to_end_with_real_soffice(tmp_path):
    doc = Document()
    numbering_el = doc.part.numbering_part.element
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    numbering_el.append(
        parse_xml(f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="95">
          <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
        </w:abstractNum>""")
    )
    numbering_el.append(
        parse_xml(f'<w:num {nsdecls("w")} w:numId="95"><w:abstractNumId w:val="95"/></w:num>')
    )

    _add_toc_field_paragraph(doc, title="Gioi thieu", page="1")
    p = doc.add_paragraph("Gioi thieu", style="Heading 1")
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = 95
    doc.add_paragraph("Noi dung.")

    path = tmp_path / "doc.docx"
    doc.save(path)

    ok = brd_toc.refresh_toc(path)
    assert ok is True

    reopened = Document(path)
    toc_paragraphs = [p for p in reopened.paragraphs if (p.style.name or "") == "toc 1"]
    assert len(toc_paragraphs) == 1
    assert "Gioi thieu" in toc_paragraphs[0].text
